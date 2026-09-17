# -*- coding: utf-8 -*-
"""内容提取脚本：读取 PPTX，输出每页的内容模型（JSON）
用于 skill 第 1 步"内容理解"。
用法: python inspect_deck.py <输入.pptx> [输出.json]
"""
import sys, json
from pptx import Presentation
from pptx.util import Emu


def cm(v):
    return round(Emu(v).cm, 2) if v is not None else None


def shape_info(sh):
    info = {
        "name": sh.name,
        "type": str(sh.shape_type),
        "x": cm(sh.left), "y": cm(sh.top),
        "w": cm(sh.width), "h": cm(sh.height),
    }
    if sh.has_text_frame and sh.text_frame.text.strip():
        info["text"] = sh.text_frame.text
        # 提取字号，用于判断层级
        sizes = set()
        for p in sh.text_frame.paragraphs:
            for r in p.runs:
                if r.font.size:
                    sizes.add(r.font.size.pt)
        if sizes:
            info["字号集合"] = sorted(sizes)
    if getattr(sh, "has_table", False) and sh.has_table:
        info["table"] = [
            [cell.text for cell in row.cells] for row in sh.table.rows
        ]
    return info


def main():
    src = sys.argv[1]
    out = sys.argv[2] if len(sys.argv) > 2 else None
    prs = Presentation(src)
    model = {
        "文件": src,
        "页数": len(prs.slides),
        "页面尺寸cm": [cm(prs.slide_width), cm(prs.slide_height)],
        "页面": [],
    }
    for i, slide in enumerate(prs.slides):
        texts = [sh for sh in slide.shapes if sh.has_text_frame and sh.text_frame.text.strip()]
        title = ""
        # action title = 最上方的非空文本（通常是标题占位符）
        if texts:
            top = min(texts, key=lambda s: (s.top or 0))
            title = top.text_frame.text[:120]
        model["页面"].append({
            "页码": i + 1,
            "版式": slide.slide_layout.name,
            "action_title": title,
            "形状数": len(slide.shapes),
            "元素": [shape_info(sh) for sh in slide.shapes],
        })
    s = json.dumps(model, ensure_ascii=False, indent=1)
    if out:
        open(out, "w", encoding="utf-8").write(s)
        print(f"已写出 {out}")
    else:
        print(s)


if __name__ == "__main__":
    main()
